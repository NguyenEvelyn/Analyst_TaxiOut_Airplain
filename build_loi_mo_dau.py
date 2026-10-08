from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Loi_mo_dau_de_tai_phan_tich_TaxiOut.docx"


def set_run_font(run, name="Times New Roman", size=13):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def add_body(doc, text, *, first_line=True, before=0, after=6):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=13)
    run.bold = True
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.left_indent = Cm(0.8)
    paragraph.paragraph_format.first_line_indent = Cm(-0.4)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run_font(run, size=11)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3.2)
section.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(13)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

title_style = doc.styles["Title"]
title_style.font.name = "Times New Roman"
title_style.font.size = Pt(18)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
title_ppr = title_style._element.get_or_add_pPr()
border = title_ppr.find(qn("w:pBdr"))
if border is not None:
    title_ppr.remove(border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(6)
title.paragraph_format.space_after = Pt(20)
run = title.add_run("LỜI MỞ ĐẦU")
set_run_font(run, size=18)
run.bold = True

add_section_heading(doc, "1. Lý do chọn đề tài và tính cấp thiết")
add_body(doc, "Trong hoạt động hàng không, TaxiOut là khoảng thời gian từ khi máy bay rời vị trí đỗ đến lúc cất cánh. Chỉ số này phản ánh một phần mức độ ùn tắc trên bề mặt sân bay và có ảnh hưởng trực tiếp đến độ đúng giờ, mức tiêu thụ nhiên liệu, chi phí khai thác và lượng phát thải. Khi lưu lượng chuyến bay gia tăng, việc nhận diện sân bay, khung giờ và điều kiện có nguy cơ TaxiOut kéo dài trở thành một nhu cầu quan trọng trong phân tích hiệu năng hàng không.")
add_body(doc, "Dữ liệu chuyến bay có quy mô lớn, nhiều thuộc tính và thay đổi theo thời gian, hãng hàng không, sân bay và mạng lưới đường bay. Việc làm sạch, tổng hợp mật độ và tạo đặc trưng từ hàng triệu bản ghi có thể vượt quá khả năng xử lý thuận tiện của các công cụ chạy trên một máy. Đây là cơ sở để lựa chọn hướng tiếp cận Big Data, trong đó HDFS hỗ trợ lưu trữ phân tán, Apache Spark thực hiện xử lý dữ liệu, H2O AutoML huấn luyện mô hình học máy và Streamlit trực quan hóa kết quả.")
add_body(doc, "Mặt khác, nhiều hệ thống dự báo chỉ công bố một chỉ số tổng thể mà chưa giải thích mô hình sai nhiều ở sân bay, giờ hoặc nhóm chuyến nào. Dữ liệu thực tế còn tồn tại các trường hợp tắc nghẽn hiếm nhưng có sai số rất lớn. Vì vậy, đề tài “Phân tích hiệu năng sân bay và dự đoán thời gian TaxiOut bằng công nghệ Big Data” được lựa chọn nhằm xây dựng một quy trình không chỉ dự đoán mà còn phân tích sâu sai số, kiểm tra độ tin cậy và xác định rõ phạm vi sử dụng của kết quả.")

add_section_heading(doc, "2. Mục tiêu nghiên cứu")
add_body(doc, "Mục tiêu chung của đề tài là xây dựng và đánh giá một hệ thống phân tích dữ liệu lớn phục vụ nhận diện quy luật ùn tắc sân bay và dự đoán thời gian TaxiOut. Các mục tiêu cụ thể gồm:", first_line=False)
add_bullet(doc, "Xây dựng quy trình lưu trữ, làm sạch, chuyển đổi và tổng hợp dữ liệu chuyến bay quy mô lớn bằng HDFS và Apache Spark.")
add_bullet(doc, "Tạo các đặc trưng phản ánh thời gian, sân bay, tuyến bay, hãng hàng không, mật độ khởi hành và lịch sử TaxiOut; đồng thời kiểm soát rò rỉ dữ liệu khi chia tập theo thời gian.")
add_bullet(doc, "Sử dụng H2O AutoML để huấn luyện, so sánh và lựa chọn mô hình hồi quy dự đoán TaxiOut; đánh giá bằng MAE, RMSE, R², bias và các mô hình đối chứng.")
add_bullet(doc, "Phân tích sai số theo nhiều chiều và xây dựng dashboard giúp người dùng quan sát tắc nghẽn, thử nghiệm dự đoán, nhận biết dữ liệu ngoài phạm vi huấn luyện và hiểu đúng giới hạn mô hình.")

add_section_heading(doc, "3. Đối tượng và phạm vi nghiên cứu")
add_body(doc, "Đối tượng nghiên cứu là thời gian TaxiOut và mối quan hệ giữa TaxiOut với các thuộc tính lịch bay, hãng hàng không, sân bay đi, sân bay đến, khoảng cách, mật độ khởi hành và thống kê vận hành lịch sử. Khách thể dữ liệu là các chuyến bay nội địa Hoa Kỳ có thông tin cần thiết cho quá trình làm sạch, tạo đặc trưng và đánh giá mô hình.")
add_bullet(doc, "Phạm vi không gian: các sân bay và đường bay nội địa Hoa Kỳ xuất hiện trong bộ dữ liệu nghiên cứu.")
add_bullet(doc, "Phạm vi thời gian: dữ liệu năm 2008; tháng 1-9 dùng để huấn luyện, tháng 10 dùng để xác thực và tháng 11-12 dùng để kiểm tra.")
add_bullet(doc, "Phạm vi dữ liệu: khoảng 6,85 triệu bản ghi được xử lý bằng Spark; thí nghiệm mô hình sử dụng 450.000 dòng train, 140.000 dòng validation và 180.000 dòng test.")
add_bullet(doc, "Phạm vi nội dung: ETL phân tán, tạo đặc trưng, hồi quy dự đoán TaxiOut, phân tích sai số, trực quan hóa và kiểm chứng nguyên mẫu; không bao gồm điều hành bay thời gian thực.")
add_body(doc, "Dữ liệu chưa bao gồm thời tiết, cấu hình đường băng, vị trí cổng, loại tàu bay, hàng đợi cất cánh và trạng thái kiểm soát không lưu. Đồng thời, dữ liệu năm 2008 không phản ánh đầy đủ hạ tầng và lịch khai thác hiện nay. Vì vậy, kết quả được giới hạn ở mục tiêu nghiên cứu phân tích và xây dựng prototype, chưa được xem là cơ sở độc lập cho quyết định vận hành.")

add_section_heading(doc, "4. Phương pháp nghiên cứu")
add_bullet(doc, "Phương pháp nghiên cứu tài liệu: tổng hợp cơ sở lý thuyết về TaxiOut, Big Data, HDFS, Spark, AutoML, hồi quy và các chỉ số đánh giá mô hình.")
add_bullet(doc, "Phương pháp định lượng và thống kê mô tả: phân tích số chuyến, TaxiOut trung bình, phân vị, mật độ khởi hành và sự khác biệt theo sân bay, giờ và nhóm TaxiOut.")
add_bullet(doc, "Phương pháp xử lý dữ liệu lớn: sử dụng Spark để kiểm tra schema, loại bản ghi không hợp lệ, tạo khung 30 phút, tính mật độ, xây dựng bảng đặc trưng và ghi dữ liệu Parquet.")
add_bullet(doc, "Phương pháp thực nghiệm học máy: chia dữ liệu theo thời gian, sử dụng H2O AutoML, so sánh với Global Mean và Origin-Hour Historical Mean, sau đó lựa chọn mô hình theo hiệu năng trên dữ liệu chưa dùng để huấn luyện.")
add_bullet(doc, "Phương pháp đánh giá và kiểm chứng: sử dụng MAE, RMSE, R², bias, phân tích sai số theo phân khúc, bootstrap theo ngày bay, kiểm thử phần mềm và đối chiếu đầu ra với metadata của mô hình.")
add_bullet(doc, "Phương pháp trực quan hóa: sử dụng Streamlit và Plotly để trình bày hiệu năng mô hình, mức ùn tắc sân bay, dự đoán thử nghiệm và các điểm yếu của mô hình.")

add_section_heading(doc, "5. Ý nghĩa khoa học và thực tiễn")
add_body(doc, "Về ý nghĩa khoa học, đề tài hệ thống hóa một quy trình kết hợp xử lý dữ liệu phân tán và học máy cho bài toán TaxiOut. Nghiên cứu làm rõ vai trò riêng của HDFS, Spark và H2O AutoML; xây dựng đặc trưng mật độ và historical encoding; đồng thời chứng minh sự cần thiết của việc đánh giá mô hình theo thời gian và theo từng phân khúc thay vì chỉ dựa vào một chỉ số tổng thể. Kết quả cũng cung cấp bằng chứng rằng mô hình có thể cải thiện so với baseline đơn giản nhưng vẫn chưa giải thích được phần lớn biến thiên TaxiOut, qua đó chỉ ra nhu cầu bổ sung dữ liệu vận hành và kiểm định ngoài thời gian.")
add_body(doc, "Về ý nghĩa thực tiễn, pipeline hỗ trợ xử lý hàng triệu bản ghi, xác định các sân bay và khung giờ có TaxiOut cao, tạo dữ liệu đặc trưng cho mô hình và trình bày kết quả qua dashboard. Mô hình H2O AutoML đạt MAE khoảng 5,84 phút trên 180.000 chuyến test và giảm MAE khoảng 18,82% so với Global Mean. Dashboard giúp người học, giảng viên hoặc nhà phân tích quan sát kết quả, thử nghiệm dự đoán và nhận cảnh báo khi dữ liệu đầu vào nằm ngoài phạm vi huấn luyện. Tuy nhiên, giá trị sử dụng hiện tại chủ yếu nằm ở phân tích, đào tạo và minh họa công nghệ; để ứng dụng thực tế cần dữ liệu mới, biến thời tiết và vận hành, cơ chế streaming, giám sát drift và tái huấn luyện định kỳ.")

add_section_heading(doc, "6. Kết cấu của báo cáo")
add_body(doc, "Báo cáo được tổ chức theo tiến trình từ cơ sở lý thuyết đến xây dựng và đánh giá hệ thống. Phần đầu giới thiệu bài toán TaxiOut và các công nghệ liên quan; phần tiếp theo trình bày dữ liệu, tiền xử lý và phương pháp xây dựng mô hình; phần thực nghiệm mô tả kiến trúc HDFS-Spark-H2O AutoML, dashboard, kết quả dự đoán và phân tích sai số; phần cuối tổng hợp kết quả, hạn chế và hướng phát triển của đề tài.")

add_page_number(section.footer.paragraphs[0])

doc.core_properties.title = "Lời mở đầu đề tài phân tích TaxiOut bằng công nghệ Big Data"
doc.core_properties.subject = "HDFS Apache Spark H2O AutoML và Streamlit"
doc.core_properties.author = "Nhóm thực hiện đề tài"
doc.core_properties.keywords = "TaxiOut, Big Data, HDFS, Spark, H2O AutoML, Streamlit"
doc.save(OUTPUT)
print(OUTPUT)
